import tkinter as tk
from tkinter import ttk, filedialog, messagebox

import pandas as pd
import matplotlib.pyplot as plt
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg


# =========================================================
# GLOBAL DATA
# =========================================================

df = None
sales = None
monthly_sales = None
product_sales = None
country_sales = None

total_revenue = 0
total_orders = 0
total_customers = 0
total_quantity = 0
average_order_value = 0


# =========================================================
# COLORS
# =========================================================

BG = "#F4F6F8"
SIDEBAR = "#18212B"
CARD = "#FFFFFF"
TEXT = "#1F2937"
MUTED = "#6B7280"
ACCENT = "#2563EB"
ACCENT_HOVER = "#1D4ED8"
BORDER = "#E5E7EB"


# =========================================================
# DATA ANALYSIS
# =========================================================

def analyze_data(file_path):

    global df
    global sales
    global monthly_sales
    global product_sales
    global country_sales

    global total_revenue
    global total_orders
    global total_customers
    global total_quantity
    global average_order_value

    # Load Excel
    df = pd.read_excel(file_path)

    # Remove duplicates
    df = df.drop_duplicates()

    # Make sure InvoiceDate is datetime
    df["InvoiceDate"] = pd.to_datetime(
        df["InvoiceDate"],
        errors="coerce"
    )

    # Detect cancelled invoices
    df["Cancelled"] = (
        df["Invoice"]
        .astype(str)
        .str.strip()
        .str.startswith("C")
    )

    # Keep valid sales
    sales = df[
        (~df["Cancelled"]) &
        (df["Quantity"] > 0) &
        (df["Price"] > 0)
    ].copy()

    # Revenue
    sales["Revenue"] = (
        sales["Quantity"] *
        sales["Price"]
    )

    # Month
    sales["Month"] = (
        sales["InvoiceDate"]
        .dt.to_period("M")
        .astype(str)
    )

    # =====================================================
    # KPIs
    # =====================================================

    total_revenue = sales["Revenue"].sum()

    total_orders = sales["Invoice"].nunique()

    total_customers = sales["Customer ID"].nunique()

    total_quantity = sales["Quantity"].sum()

    if total_orders > 0:
        average_order_value = (
            total_revenue /
            total_orders
        )
    else:
        average_order_value = 0

    # =====================================================
    # MONTHLY SALES
    # =====================================================

    monthly_sales = (
        sales.groupby("Month")
        .agg(
            Revenue=("Revenue", "sum"),
            Orders=("Invoice", "nunique"),
            Quantity=("Quantity", "sum")
        )
        .sort_index()
    )

    # =====================================================
    # PRODUCT SALES
    # =====================================================

    product_sales = (
        sales.dropna(subset=["Description"])
        .groupby("Description")
        .agg(
            Revenue=("Revenue", "sum"),
            Quantity=("Quantity", "sum"),
            Orders=("Invoice", "nunique")
        )
        .sort_values(
            "Revenue",
            ascending=False
        )
    )

    # =====================================================
    # COUNTRY SALES
    # =====================================================

    country_sales = (
        sales.dropna(subset=["Country"])
        .groupby("Country")
        .agg(
            Revenue=("Revenue", "sum"),
            Orders=("Invoice", "nunique"),
            Customers=("Customer ID", "nunique")
        )
        .sort_values(
            "Revenue",
            ascending=False
        )
    )


# =========================================================
# FILTERS
# =========================================================

selected_sales = None


def apply_filters():

    global selected_sales

    if not data_loaded():
        return

    selected_sales = sales.copy()

    # Country filter
    country = country_var.get()

    if country != "All Countries":
        selected_sales = selected_sales[
            selected_sales["Country"] == country
        ]

    # Start date
    start_date = start_date_entry.get().strip()

    if start_date:
        try:
            start_date = pd.to_datetime(start_date)

            selected_sales = selected_sales[
                selected_sales["InvoiceDate"] >= start_date
            ]

        except:
            messagebox.showerror(
                "Invalid Date",
                "Please enter the start date correctly.\n\n"
                "Example: 2011-01-01"
            )
            return

    # End date
    end_date = end_date_entry.get().strip()

    if end_date:
        try:
            end_date = pd.to_datetime(end_date)

            selected_sales = selected_sales[
                selected_sales["InvoiceDate"] <= end_date
            ]

        except:
            messagebox.showerror(
                "Invalid Date",
                "Please enter the end date correctly.\n\n"
                "Example: 2011-12-31"
            )
            return

    update_filtered_kpis()

    show_monthly_revenue()


def clear_filters():

    country_var.set("All Countries")

    start_date_entry.delete(
        0,
        tk.END
    )

    end_date_entry.delete(
        0,
        tk.END
    )

    apply_filters()


def update_filtered_kpis():

    global selected_sales

    if selected_sales is None:
        return

    filtered_revenue = selected_sales["Revenue"].sum()

    filtered_orders = selected_sales["Invoice"].nunique()

    filtered_customers = selected_sales["Customer ID"].nunique()

    if filtered_orders > 0:
        filtered_aov = (
            filtered_revenue /
            filtered_orders
        )
    else:
        filtered_aov = 0

    revenue_value.config(
        text=f"{filtered_revenue:,.0f}"
    )

    orders_value.config(
        text=f"{filtered_orders:,}"
    )

    customers_value.config(
        text=f"{filtered_customers:,}"
    )

    aov_value.config(
        text=f"{filtered_aov:,.2f}"
    )

# =========================================================
# CHECK DATA
# =========================================================

def data_loaded():

    if df is None:

        messagebox.showwarning(
            "No Data",
            "Please upload an Excel file first."
        )

        return False

    return True
# =========================================================
# UPLOAD FILE
# =========================================================

def upload_file():

    global selected_sales
    file_path = filedialog.askopenfilename(
        title="Select Excel File",
        filetypes=[
            ("Excel Files", "*.xlsx *.xls"),
            ("All Files", "*.*")
        ]
    )

    if not file_path:
        return

    try:

        analyze_data(file_path)
        country_dropdown["values"] = (
    ["All Countries"] +
    sorted(
        sales["Country"]
        .dropna()
        .unique()
        .tolist()
    )
)

        country_var.set("All Countries")
        selected_sales = sales.copy()

        file_name = file_path.split("/")[-1]

        file_label.config(
            text=f"Loaded: {file_name}"
        )

        status_label.config(
            text="Analysis ready",
            foreground="#16A34A"
        )

        update_kpis()

        show_dashboard()

    except Exception as e:

        status_label.config(
            text="Error loading file",
            foreground="#DC2626"
        )

        messagebox.showerror(
            "Analysis Error",
            f"Could not analyze the selected file.\n\n{e}"
        )


# =========================================================
# UPDATE KPI CARDS
# =========================================================

def update_kpis():

    revenue_value.config(
        text=f"{total_revenue:,.0f}"
    )

    orders_value.config(
        text=f"{total_orders:,}"
    )

    customers_value.config(
        text=f"{total_customers:,}"
    )

    aov_value.config(
        text=f"{average_order_value:,.2f}"
    )


# =========================================================
# CLEAR CHART
# =========================================================

def clear_chart():

    for widget in chart_frame.winfo_children():
        widget.destroy()


# =========================================================
# SHOW DASHBOARD
# =========================================================

def show_dashboard():

    if not data_loaded():
        return

    clear_chart()

    figure = plt.Figure(
        figsize=(9, 5.2),
        dpi=100
    )

    ax = figure.add_subplot(111)

    ax.text(
        0.5,
        0.5,
        "Select a report from the left",
        ha="center",
        va="center",
        fontsize=16,
        color=MUTED
    )

    ax.axis("off")

    canvas = FigureCanvasTkAgg(
        figure,
        master=chart_frame
    )

    canvas.draw()

    canvas.get_tk_widget().pack(
        fill="both",
        expand=True,
        padx=15,
        pady=15
    )


# =========================================================
# MONTHLY REVENUE
# =========================================================
def show_monthly_revenue():

    if not data_loaded():
        return

    clear_chart()

    figure = plt.Figure(
        figsize=(9, 5.2),
        dpi=100
    )

    ax = figure.add_subplot(111)

    ax.plot(
        monthly_sales.index,
        monthly_sales["Revenue"],
        marker="o",
        linewidth=2
    )

    ax.set_title(
        "Monthly Revenue",
        fontsize=16,
        fontweight="bold",
        pad=15
    )

    ax.set_xlabel("Month")

    ax.set_ylabel("Revenue")

    ax.tick_params(
        axis="x",
        rotation=45
    )

    ax.grid(
        axis="y",
        alpha=0.25
    )

    figure.tight_layout(rect=[0,0,1,0.93])

    canvas = FigureCanvasTkAgg(
        figure,
        master=chart_frame
    )

    canvas.draw()

    canvas.get_tk_widget().pack(
        fill="both",
        expand=True,
        padx=15,
        pady=15
    )


# =========================================================
# TOP PRODUCTS
# =========================================================

def show_top_products():

    if not data_loaded():
        return

    if selected_sales is None or selected_sales.empty:
        return

    clear_chart()

    filtered_product_sales = (
        selected_sales
        .dropna(subset=["Description"])
        .groupby("Description")
        .agg(
            Revenue=("Revenue", "sum"),
            Quantity=("Quantity", "sum"),
            Orders=("Invoice", "nunique")
        )
        .sort_values(
            "Revenue",
            ascending=False
        )
    )

    top_products = (
        filtered_product_sales
        .head(10)
        .sort_values("Revenue")
    )

    figure = plt.Figure(
        figsize=(9, 5.2),
        dpi=100
    )

    ax = figure.add_subplot(111)

    ax.barh(
        top_products.index,
        top_products["Revenue"]
    )

    ax.set_title(
        "Top 10 Products by Revenue",
        fontsize=16,
        fontweight="bold",
        pad=15
    )

    ax.set_xlabel("Revenue")

    ax.grid(
        axis="x",
        alpha=0.25
    )

    figure.tight_layout(rect=[0,0,1,0.93])

    canvas = FigureCanvasTkAgg(
        figure,
        master=chart_frame
    )

    canvas.draw()

    canvas.get_tk_widget().pack(
        fill="both",
        expand=True,
        padx=15,
        pady=15
    )


# =========================================================
# TOP COUNTRIES
# =========================================================

def show_top_countries():

    if not data_loaded():
        return

    if selected_sales is None or selected_sales.empty:
        return

    clear_chart()

    filtered_country_sales = (
        selected_sales
        .dropna(subset=["Country"])
        .groupby("Country")
        .agg(
            Revenue=("Revenue", "sum"),
            Orders=("Invoice", "nunique"),
            Customers=("Customer ID", "nunique")
        )
        .sort_values(
            "Revenue",
            ascending=False
        )
    )

    top_countries = (
        filtered_country_sales
        .head(10)
        .sort_values("Revenue")
    )

    figure = plt.Figure(
        figsize=(9, 5.2),
        dpi=100
    )

    ax = figure.add_subplot(111)

    ax.barh(
        top_countries.index,
        top_countries["Revenue"]
    )

    ax.set_title(
        "Top 10 Countries by Revenue",
        fontsize=16,
        fontweight="bold",
        pad=15
    )

    ax.set_xlabel("Revenue")

    ax.grid(
        axis="x",
        alpha=0.25
    )

    figure.tight_layout(rect=[0,0,1,0.93])

    canvas = FigureCanvasTkAgg(
        figure,
        master=chart_frame
    )

    canvas.draw()

    canvas.get_tk_widget().pack(
        fill="both",
        expand=True,
        padx=15,
        pady=15
    )


# =========================================================
# EXPORT FILTERED DATA
# =========================================================

def export_filtered_data():

    if not data_loaded():
        return

    if selected_sales is None or selected_sales.empty:
        messagebox.showwarning(
            "No Data",
            "There is no filtered data to export."
        )
        return

    file_path = filedialog.asksaveasfilename(
        title="Save Excel Report",
        defaultextension=".xlsx",
        filetypes=[
            ("Excel files", "*.xlsx")
        ],
        initialfile="sales_report.xlsx"
    )

    if not file_path:
        return

    try:

        from openpyxl import Workbook
        from openpyxl.styles import (
            Font,
            PatternFill,
            Border,
            Side,
            Alignment
        )
        from openpyxl.chart import (
            LineChart,
            BarChart,
            Reference
        )
        from openpyxl.chart.label import DataLabelList
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo

        filtered_data = selected_sales.copy()

        # =====================================================
        # PREPARE REPORT DATA
        # =====================================================

        filtered_monthly = (
            filtered_data
            .assign(
                Month=filtered_data["InvoiceDate"]
                .dt.to_period("M")
                .astype(str)
            )
            .groupby("Month", as_index=False)["Revenue"]
            .sum()
            .sort_values("Month")
        )

        filtered_products = (
            filtered_data
            .dropna(subset=["Description"])
            .groupby("Description", as_index=False)
            .agg(
                Revenue=("Revenue", "sum"),
                Quantity=("Quantity", "sum"),
                Orders=("Invoice", "nunique")
            )
            .sort_values(
                "Revenue",
                ascending=False
            )
        )

        top_products = filtered_products.head(10).copy()

        filtered_countries = (
            filtered_data
            .dropna(subset=["Country"])
            .groupby("Country", as_index=False)
            .agg(
                Revenue=("Revenue", "sum"),
                Orders=("Invoice", "nunique"),
                Customers=("Customer ID", "nunique")
            )
            .sort_values(
                "Revenue",
                ascending=False
            )
        )

        top_countries = filtered_countries.head(10).copy()

        # =====================================================
        # KPI
        # =====================================================

        report_revenue = filtered_data["Revenue"].sum()

        report_orders = filtered_data["Invoice"].nunique()

        report_customers = (
            filtered_data["Customer ID"]
            .nunique()
        )

        report_quantity = (
            filtered_data["Quantity"]
            .sum()
        )

        report_aov = (
            report_revenue / report_orders
            if report_orders > 0
            else 0
        )

        # =====================================================
        # MANAGEMENT INSIGHTS
        # =====================================================

        top_product_name = (
            top_products.iloc[0]["Description"]
            if not top_products.empty
            else "N/A"
        )

        top_product_revenue = (
            top_products.iloc[0]["Revenue"]
            if not top_products.empty
            else 0
        )

        top_country_name = (
            top_countries.iloc[0]["Country"]
            if not top_countries.empty
            else "N/A"
        )

        top_country_revenue = (
            top_countries.iloc[0]["Revenue"]
            if not top_countries.empty
            else 0
        )

        best_month = (filtered_monthly.sort_values("Revenue",ascending=False).iloc[0]["Month"]
                      if not filtered_monthly.empty
                      else "N/A")

        top_10_revenue = (
            top_products["Revenue"].sum()
            if not top_products.empty
            else 0
        )

        top_10_share = (
            (top_10_revenue / report_revenue) * 100
            if report_revenue > 0
            else 0
        )

        # =====================================================
        # WORKBOOK
        # =====================================================

        workbook = Workbook()

        summary_sheet = workbook.active
        summary_sheet.title = "Summary"

        monthly_sheet = workbook.create_sheet(
            "Monthly Revenue"
        )

        products_sheet = workbook.create_sheet(
            "Top Products"
        )

        filtered_sheet = workbook.create_sheet(
            "Filtered Data"
        )

        # =====================================================
        # STYLES
        # =====================================================

        dark_fill = PatternFill(
            "solid",
            fgColor="18212B"
        )

        blue_fill = PatternFill(
            "solid",
            fgColor="2563EB"
        )

        light_fill = PatternFill(
            "solid",
            fgColor="EAF2FF"
        )

        gray_fill = PatternFill(
            "solid",
            fgColor="F3F4F6"
        )

        white_font = Font(
            color="FFFFFF",
            bold=True
        )

        title_font = Font(
            size=20,
            bold=True,
            color="FFFFFF"
        )

        section_font = Font(
            size=13,
            bold=True,
            color="18212B"
        )

        normal_font = Font(
            size=11,
            color="1F2937"
        )

        thin_border = Border(
            bottom=Side(
                style="thin",
                color="D1D5DB"
            )
        )

        center = Alignment(
            horizontal="center",
            vertical="center"
        )

        # =====================================================
        # SUMMARY
        # =====================================================

        summary_sheet.merge_cells(
            "A1:F2"
        )

        summary_sheet["A1"] = (
            "SALES ANALYTICS REPORT"
        )

        summary_sheet["A1"].fill = dark_fill
        summary_sheet["A1"].font = title_font
        summary_sheet["A1"].alignment = center

        summary_sheet["A4"] = "Report Information"
        summary_sheet["A4"].font = section_font

        summary_sheet["A5"] = "Generated"
        summary_sheet["B5"] = pd.Timestamp.now().strftime(
            "%Y-%m-%d %H:%M"
        )

        summary_sheet["A6"] = "Rows Analyzed"
        summary_sheet["B6"] = len(filtered_data)

        summary_sheet["A7"] = "Date From"
        summary_sheet["B7"] = (
            filtered_data["InvoiceDate"].min()
            .strftime("%Y-%m-%d")
        )

        summary_sheet["A8"] = "Date To"
        summary_sheet["B8"] = (
            filtered_data["InvoiceDate"].max()
            .strftime("%Y-%m-%d")
        )

        summary_sheet["A10"] = "Key Performance Indicators"
        summary_sheet["A10"].font = section_font

        kpis = [
            ("Total Revenue", report_revenue),
            ("Total Orders", report_orders),
            ("Total Customers", report_customers),
            ("Total Quantity", report_quantity),
            ("Average Order Value", report_aov)
        ]

        row = 11

        for name, value in kpis:

            summary_sheet.cell(
                row=row,
                column=1,
                value=name
            )

            summary_sheet.cell(
                row=row,
                column=2,
                value=value
            )

            summary_sheet.cell(
                row=row,
                column=1
            ).fill = light_fill

            summary_sheet.cell(
                row=row,
                column=1
            ).font = Font(
                bold=True
            )

            row += 1

        # =====================================================
        # MANAGEMENT INSIGHTS
        # =====================================================

        summary_sheet["D4"] = "Management Insights"
        summary_sheet["D4"].font = section_font

        insights = [
            (
                "Top Product",
                top_product_name
            ),
            (
                "Top Product Revenue",
                top_product_revenue
            ),
            (
                "Top Country",
                top_country_name
            ),
            (
                "Top Country Revenue",
                top_country_revenue
            ),
            (
                "Best Month",
                best_month
            ),
            (
                "Top 10 Products Share",
                top_10_share / 100
            )
        ]

        row = 5

        for name, value in insights:

            summary_sheet.cell(
                row=row,
                column=4,
                value=name
            )

            summary_sheet.cell(
                row=row,
                column=5,
                value=value
            )

            summary_sheet.cell(
                row=row,
                column=4
            ).fill = light_fill

            summary_sheet.cell(
                row=row,
                column=4
            ).font = Font(
                bold=True
            )

            row += 1

        # =====================================================
        # SUMMARY FORMATTING
        # =====================================================

        for cell in summary_sheet["A"]:
            cell.font = cell.font.copy(
                name="Arial"
            )

        for row_cells in summary_sheet.iter_rows():

            for cell in row_cells:

                cell.alignment = Alignment(
                    vertical="center"
                )

        summary_sheet["B11"].number_format = '#,##0.00'
        summary_sheet["B12"].number_format = '#,##0'
        summary_sheet["B13"].number_format = '#,##0'
        summary_sheet["B14"].number_format = '#,##0'
        summary_sheet["B15"].number_format = '#,##0.00'

        summary_sheet["E6"].number_format = '#,##0.00'
        summary_sheet["E8"].number_format = '#,##0.00'
        summary_sheet["E10"].number_format = '0.0%'

        # =====================================================
        # MONTHLY REVENUE
        # =====================================================

        monthly_sheet.merge_cells("A1:D2")

        monthly_sheet["A1"] = "MONTHLY REVENUE"

        monthly_sheet["A1"].fill = dark_fill
        monthly_sheet["A1"].font = title_font
        monthly_sheet["A1"].alignment = center

        monthly_sheet["A4"] = "Month"
        monthly_sheet["B4"] = "Revenue"

        for cell in monthly_sheet[4]:
            cell.fill = blue_fill
            cell.font = white_font
            cell.alignment = center

        for index, data_row in enumerate(
            filtered_monthly.itertuples(index=False),
            start=5
        ):

            monthly_sheet.cell(
                index,
                1,
                data_row.Month
            )

            monthly_sheet.cell(
                index,
                2,
                data_row.Revenue
            )

            monthly_sheet.cell(
                index,
                2
            ).number_format = '#,##0.00'

        # =====================================================
        # MONTHLY CHART
        # =====================================================

        if len(filtered_monthly) > 0:

            line_chart = LineChart()

            line_chart.title = (
                "Monthly Revenue Trend"
            )

            line_chart.y_axis.title = "Revenue"
            line_chart.x_axis.title = "Month"

            line_chart.height = 8
            line_chart.width = 16

            data = Reference(
                monthly_sheet,
                min_col=2,
                min_row=4,
                max_row=4 + len(filtered_monthly)
            )
            categories = Reference(
                monthly_sheet,
                min_col=1,
                min_row=5,
                max_row=4 + len(filtered_monthly)
            )

            line_chart.add_data(
                data,
                titles_from_data=True
            )

            line_chart.set_categories(
                categories
            )

            line_chart.legend = None

            monthly_sheet.add_chart(
                line_chart,
                "D4"
            )

        # =====================================================
        # TOP PRODUCTS
        # =====================================================

        products_sheet.merge_cells("A1:E2")

        products_sheet["A1"] = (
            "TOP 10 PRODUCTS"
        )

        products_sheet["A1"].fill = dark_fill
        products_sheet["A1"].font = title_font
        products_sheet["A1"].alignment = center

        headers = [
            "Product",
            "Revenue",
            "Quantity",
            "Orders"
        ]

        for col, header in enumerate(
            headers,
            start=1
        ):

            cell = products_sheet.cell(
                4,
                col,
                header
            )

            cell.fill = blue_fill
            cell.font = white_font
            cell.alignment = center

        for index, data_row in enumerate(
            top_products.itertuples(index=False),
            start=5
        ):

            products_sheet.cell(
                index,
                1,
                data_row.Description
            )

            products_sheet.cell(
                index,
                2,
                data_row.Revenue
            )

            products_sheet.cell(
                index,
                3,
                data_row.Quantity
            )

            products_sheet.cell(
                index,
                4,
                data_row.Orders
            )

            products_sheet.cell(
                index,
                2
            ).number_format = '#,##0.00'

        # =====================================================
        # PRODUCTS CHART
        # =====================================================

        if len(top_products) > 0:

            bar_chart = BarChart()

            bar_chart.type = "bar"

            bar_chart.style = 10

            bar_chart.title = (
                "Top 10 Products by Revenue"
            )

            bar_chart.y_axis.title = "Product"
            bar_chart.x_axis.title = "Revenue"

            bar_chart.height = 10
            bar_chart.width = 18

            data = Reference(
                products_sheet,
                min_col=2,
                min_row=4,
                max_row=4 + len(top_products)
            )

            categories = Reference(
                products_sheet,
                min_col=1,
                min_row=5,
                max_row=4 + len(top_products)
            )

            bar_chart.add_data(
                data,
                titles_from_data=True
            )

            bar_chart.set_categories(
                categories
            )

            bar_chart.legend = None

            bar_chart.dataLabels = DataLabelList()
            bar_chart.dataLabels.showVal = True

            products_sheet.add_chart(
                bar_chart,
                "F4"
            )

        # =====================================================
        # FILTERED DATA
        # =====================================================

        for col_index, column_name in enumerate(
            filtered_data.columns,
            start=1
        ):

            cell = filtered_sheet.cell(
                1,
                col_index,
                column_name
            )

            cell.fill = dark_fill
            cell.font = white_font
            cell.alignment = center

        for row_index, row_data in enumerate(
            filtered_data.itertuples(index=False),
            start=2
        ):

            for col_index, value in enumerate(
                row_data,
                start=1
            ):

                cell = filtered_sheet.cell(
                    row_index,
                    col_index,
                    value
                )

                if (
                    filtered_data.columns[col_index - 1]
                    == "InvoiceDate"
                    and pd.notna(value)
                ):
                    cell.number_format = (
                        "yyyy-mm-dd hh:mm"
                    )

                if (
                    filtered_data.columns[col_index - 1]
                    == "Revenue"
                ):
                    cell.number_format = (
                        '#,##0.00'
                    )

        # =====================================================
        # FILTERED DATA TABLE
        # =====================================================

        if len(filtered_data) > 0:

            last_row = len(filtered_data) + 1
            last_col = len(filtered_data.columns)

            table_ref = (
                f"A1:"
                f"{get_column_letter(last_col)}"
                f"{last_row}"
            )

            table = Table(
                displayName="FilteredSalesData",
                ref=table_ref
            )

            table_style = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False
            )

            table.tableStyleInfo = table_style

            filtered_sheet.add_table(
                table
            )

        filtered_sheet.freeze_panes = "A2"

        # =====================================================
        # COLUMN WIDTHS
        # =====================================================

        for sheet in workbook.worksheets:

            for column_cells in sheet.columns:

                max_length = 0

                column_letter = (
                    get_column_letter(
                        column_cells[0].column
                    )
                )

                for cell in column_cells:

                    try:

                        cell_length = len(
                            str(cell.value)
                        )

                        if cell_length > max_length:
                            max_length = cell_length

                    except:
                        pass

                sheet.column_dimensions[
                    column_letter
                ].width = min(
                    max(max_length + 2, 12),
                    45
                )

        # =====================================================
        # FREEZE PANES
        # =====================================================

        monthly_sheet.freeze_panes = "A5"
        products_sheet.freeze_panes = "A5"

        # =====================================================
        # SAVE
        # =====================================================

        workbook.save(file_path)

        messagebox.showinfo(
            "Export Successful",
            "Professional Excel report created successfully."
        )

    except Exception as e:

        messagebox.showerror(
            "Export Error",
            f"An error occurred while creating the report:\n\n{e}"
        )
# =========================================================
# MAIN WINDOW
# =========================================================

window = tk.Tk()

window.title(
    "Online Retail Sales Analytics"
)

window.geometry(
    "1250x750"
)

window.minsize(
    1050,
    650
)

window.configure(
    bg=BG
)


# =========================================================
# SIDEBAR
# =========================================================

sidebar = tk.Frame(
    window,
    bg=SIDEBAR,
    width=230
)

sidebar.pack(
    side="left",
    fill="y"
)

sidebar.pack_propagate(False)


# Logo / title
logo_label = tk.Label(
    sidebar,
    text="RETAIL\nANALYTICS",
    bg=SIDEBAR,
    fg="white",
    font=("Arial", 20, "bold"),
    justify="left"
)

logo_label.pack(
    anchor="w",
    padx=25,
    pady=(30, 10)
)


sidebar_subtitle = tk.Label(
    sidebar,
    text="Sales Intelligence",
    bg=SIDEBAR,
    fg="#9CA3AF",
    font=("Arial", 10)
)

sidebar_subtitle.pack(
    anchor="w",
    padx=27,
    pady=(0, 30)
)


# Upload button
upload_button = tk.Button(
    sidebar,
    text="  Upload Excel File",
    command=upload_file,
    bg=ACCENT,
    fg="white",
    activebackground=ACCENT_HOVER,
    activeforeground="white",
    relief="flat",
    bd=0,
    font=("Arial", 11, "bold"),
    anchor="w",
    padx=15,
    pady=12,
    cursor="hand2"
)

upload_button.pack(
    fill="x",
    padx=15,
    pady=5
)


# Reports label
reports_title = tk.Label(
    sidebar,
    text="REPORTS",
    bg=SIDEBAR,
    fg="#6B7280",
    font=("Arial", 9, "bold")
)

reports_title.pack(
    anchor="w",
    padx=25,
    pady=(30, 10)
)


# Monthly Revenue button
monthly_button = tk.Button(
    sidebar,
    text="  Monthly Revenue",
    command=show_monthly_revenue,
    bg=SIDEBAR,
    fg="#D1D5DB",
    activebackground="#263442",
    activeforeground="white",
    relief="flat",
    bd=0,
    font=("Arial", 10),
    anchor="w",
    padx=15,
    pady=11,
    cursor="hand2"
)

monthly_button.pack(
    fill="x",
    padx=15
)


# Products button
products_button = tk.Button(sidebar,text="  Top 10 Products",command=show_top_products,bg=SIDEBAR,fg="#D1D5DB",activebackground="#263442",activeforeground="white",relief="flat",bd=0,font=("Arial", 10),anchor="w",padx=15,pady=11,cursor="hand2")

products_button.pack(
    fill="x",
    padx=15
)


# Countries button
countries_button = tk.Button(
    sidebar,
    text="  Top 10 Countries",
    command=show_top_countries,
    bg=SIDEBAR,
    fg="#D1D5DB",
    activebackground="#263442",
    activeforeground="white",
    relief="flat",
    bd=0,
    font=("Arial", 10),
    anchor="w",
    padx=15,
    pady=11,
    cursor="hand2"
)

countries_button.pack(
    fill="x",
    padx=15
)
export_button = tk.Button(
    sidebar,
    text="  Export Excel Report",
    command=export_filtered_data,
    bg="#16A34A",
    fg="white",
    activebackground="#15803D",
    activeforeground="white",
    relief="flat",
    bd=0,
    font=("Arial", 10, "bold"),
    anchor="w",
    padx=15,
    pady=11,
    cursor="hand2"
)

export_button.pack(
    fill="x",
    padx=15,
    pady=(20, 0)
)


# =========================================================
# MAIN CONTENT
# =========================================================
content = tk.Frame(
    window,
    bg=BG
)

content.pack(
    side="left",
    fill="both",
    expand=True
)


# =========================================================
# HEADER
# =========================================================

header = tk.Frame(
    content,
    bg=BG
)

header.pack(
    fill="x",
    padx=30,
    pady=(25, 10)
)


header_title = tk.Label(
    header,
    text="Sales Dashboard",
    bg=BG,
    fg=TEXT,
    font=("Arial", 24, "bold")
)

header_title.pack(
    side="left"
)


status_label = tk.Label(
    header,
    text="No file loaded",
    bg=BG,
    fg=MUTED,
    font=("Arial", 10)
)

status_label.pack(
    side="right"
)


# =========================================================
# FILE INFORMATION
# =========================================================

file_label = tk.Label(
    content,
    text="Upload an Excel file to start analysis",
    bg=BG,
    fg=MUTED,
    font=("Arial", 10)
)

file_label.pack(
    anchor="w",
    padx=32,
    pady=(0, 15)
)
# =========================================================
# FILTER PANEL
# =========================================================

filter_frame = tk.Frame(
    content,
    bg=CARD,
    highlightbackground=BORDER,
    highlightthickness=1
)

filter_frame.pack(
    fill="x",
    padx=32,
    pady=(0, 15)
)


filter_title = tk.Label(
    filter_frame,
    text="Filters",
    bg=CARD,
    fg=TEXT,
    font=("Arial", 12, "bold")
)

filter_title.grid(
    row=0,
    column=0,
    padx=15,
    pady=(12, 5),
    sticky="w"
)


# Country
country_label = tk.Label(
    filter_frame,
    text="Country",
    bg=CARD,
    fg=MUTED,
    font=("Arial", 9)
)

country_label.grid(
    row=1,
    column=0,
    padx=15,
    pady=5,
    sticky="w"
)


country_var = tk.StringVar(
    value="All Countries"
)


country_dropdown = ttk.Combobox(
    filter_frame,
    textvariable=country_var,
    state="readonly",
    width=20
)

country_dropdown.grid(
    row=2,
    column=0,
    padx=15,
    pady=(0, 15)
)


# Start Date
start_date_label = tk.Label(
    filter_frame,
    text="Start Date",
    bg=CARD,
    fg=MUTED,
    font=("Arial", 9)
)

start_date_label.grid(
    row=1,
    column=1,
    padx=15,
    pady=5,
    sticky="w"
)


start_date_entry = tk.Entry(
    filter_frame,
    width=18
)

start_date_entry.grid(
    row=2,
    column=1,
    padx=15,
    pady=(0, 15)
)


# End Date
end_date_label = tk.Label(
    filter_frame,
    text="End Date",
    bg=CARD,
    fg=MUTED,
    font=("Arial", 9)
)

end_date_label.grid(
    row=1,
    column=2,
    padx=15,
    pady=5,
    sticky="w"
)


end_date_entry = tk.Entry(
    filter_frame,
    width=18
)

end_date_entry.grid(
    row=2,
    column=2,
    padx=15,
    pady=(0, 15)
)


# Apply
apply_button = tk.Button(
    filter_frame,
    text="Apply Filters",
    command=apply_filters,
    bg=ACCENT,
    fg="white",
    activebackground=ACCENT_HOVER,
    relief="flat",
    bd=0,
    font=("Arial", 10, "bold"),
    padx=15,
    pady=7,
    cursor="hand2"
)

apply_button.grid(
    row=2,
    column=3,
    padx=10
)


# Clear
clear_button = tk.Button(
    filter_frame,
    text="Clear",
    command=clear_filters,
    bg="#E5E7EB",
    fg=TEXT,
    relief="flat",
    bd=0,
    font=("Arial", 10),
    padx=15,
    pady=7,
    cursor="hand2"
)

clear_button.grid(
    row=2,
    column=4,
    padx=10
)


# =========================================================
# KPI CARDS
# =========================================================

kpi_frame = tk.Frame(
    content,
    bg=BG
)

kpi_frame.pack(
    fill="x",
    padx=25,
    pady=5
)


def create_kpi_card(parent, title):

    card = tk.Frame(
        parent,
        bg=CARD,
        highlightbackground=BORDER,
        highlightthickness=1
    )

    card.pack(
        side="left",
        fill="both",
        expand=True,
        padx=7
    )

    title_label = tk.Label(
        card,
        text=title,
        bg=CARD,
        fg=MUTED,
        font=("Arial", 10)
    )

    title_label.pack(
        anchor="w",
        padx=18,
        pady=(15, 5)
    )

    value_label = tk.Label(
        card,
        text="0",
        bg=CARD,
        fg=TEXT,
        font=("Arial", 20, "bold")
    )

    value_label.pack(
        anchor="w",
        padx=18,
        pady=(0, 15)
    )

    return value_label


revenue_value = create_kpi_card(
    kpi_frame,
    "Total Revenue"
)

orders_value = create_kpi_card(
    kpi_frame,
    "Total Orders"
)

customers_value = create_kpi_card(
    kpi_frame,
    "Total Customers"
)

aov_value = create_kpi_card(
    kpi_frame,
    "Average Order Value"
)


# =========================================================
# CHART AREA
# =========================================================
chart_card = tk.Frame(
    content,
    bg=CARD,
    highlightbackground=BORDER,
    highlightthickness=1
)

chart_card.pack(
    fill="both",
    expand=True,
    padx=32,
    pady=(15, 25)
)


chart_header = tk.Frame(
    chart_card,
    bg=CARD
)

chart_header.pack(
    fill="x",
    padx=20,
    pady=(15, 0)
)


chart_title = tk.Label(
    chart_header,
    text="Analytics",
    bg=CARD,
    fg=TEXT,
    font=("Arial", 14, "bold")
)

chart_title.pack(
    side="left"
)


chart_frame = tk.Frame(
    chart_card,
    bg=CARD
)

chart_frame.pack(
    fill="both",
    expand=True
)


# =========================================================
# START
# =========================================================

window.mainloop()